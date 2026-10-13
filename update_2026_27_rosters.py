#!/usr/bin/env python3
"""Replace site rosters and roster-backed caches with the 2026-27 LBPRC lists."""

from __future__ import annotations

import copy
import hashlib
import json
import re
import unicodedata
from collections import defaultdict
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
SANTURCE_BOOK = Path("/Users/colbymorris/Downloads/2026 Santurce League Roster-1.xlsx")
LBPRC_BOOK = Path("/Users/colbymorris/Downloads/2026-27 LBPRC PRESEASON ROSTERS-1.xlsx")

TEAM_META = {
    "SAN": (685, "Cangrejeros de Santurce"),
    "CAG": (686, "Criollos de Caguas"),
    "CAR": (687, "Gigantes de Carolina"),
    "MAY": (688, "Indios de Mayaguez"),
    "PON": (689, "Leones de Ponce"),
    "SJU": (4087, "Senadores de San Juan"),
}
SHEET_TEAM = {"CAG": "CAG", "CAR": "CAR", "MAY": "MAY", "PON": "PON", "SJ": "SJU"}
PITCHER_POS = {"P", "RHP", "LHP", "RHR", "LHR", "RHS", "LHS"}
POSITION_NAMES = {
    "PITCHER": "P",
    "CATCHER": "C",
    "OUTFIELDER": "OF",
    "SHORTSTOP": "SS",
    "SECOND BASE": "2B",
    "THIRD BASE": "3B",
    "FIRST BASE": "1B",
    "DESIGNATED HITTER": "DH",
}
TRANSACTIONS = [
    {
        "tweet_id": "2108744698424947150",
        "url": "https://x.com/mop_knows/status/2108744698424947150",
        "exact_text": (
            "Hearing that the most outstanding player in NPB Ni-Gun this year, "
            "OF Yuta Nakami(20 HR/20 D/.909 OPS), along with rookie training player "
            "Keishin Saito will be active this winter with Senadores de San Juan.\n\n"
            "Second straight year that the Lions send talent to this team.\n\n#LBPRC🇵🇷❄️"
        ),
        "applied": ["Add Yuta Nakami (OF) to SJU", "Add Keishin Saito (P) to SJU"],
    },
    {
        "tweet_id": "2108748202379727101",
        "url": "https://x.com/mop_knows/status/2108748202379727101",
        "exact_text": (
            "20-year old righty rookie Haruse Narita is also going along with them to San Juan.\n\n"
            "He posted a 2.45 ERA with 27 K’s in 29.1 innings on the farm. Has to work "
            "with his control, as he gave 20 walks in them."
        ),
        "applied": ["Add Haruse Narita (P) to SJU"],
    },
    {
        "tweet_id": "2108571420041138608",
        "url": "https://x.com/lbprc/status/2108571420041138608",
        "exact_text": (
            "MOVIMIENTOS ✅\n"
            "TRANSACCIONES — 8 DE OCTUBRE DE 2026\n"
            "Leones de Ponce — RECIBEN: +MARIO ZABALA, RHP\n"
            "Senadores de San Juan — RECIBEN: +FREDDY CABRERA, RHP"
        ),
        "applied": [
            "Add Mario Zabala (P) to PON and remove him from every other team",
            "Add Freddy Cabrera (P) to SJU and remove him from every other team",
        ],
    },
]


def load(name: str):
    return json.loads((DATA / name).read_text())


def write(name: str, value, *, compact: bool = False):
    kwargs = {"ensure_ascii": False}
    if compact:
        kwargs["separators"] = (",", ":")
    else:
        kwargs["indent"] = 2
    (DATA / name).write_text(json.dumps(value, **kwargs) + ("" if compact else "\n"))


def norm(value: str) -> str:
    value = unicodedata.normalize("NFKD", value or "")
    value = "".join(c for c in value if not unicodedata.combining(c))
    value = re.sub(r'\s+"[^"]+"', "", value)
    value = re.sub(r"\b(jr|sr|ii|iii|iv)\b", "", value, flags=re.I)
    return re.sub(r"[^a-z0-9]", "", value.lower())


# Normalize the hand-entered reserve lookup after norm() is defined.
SANTURCE_RESERVE_POSITIONS = {norm(name): pos for name, pos in {
    "Edwin Alicea": "P",
    "Luis Calderon": "P",
    "Jorhan Castro": "IF",
    "Edgar Colon": "P",
    "Dominic Hamel": "P",
    "Michael Martinez": "P",
    "Jatniel McCloud": "OF",
    "Joseph Montalvo": "P",
    "Sebastian Rolon": "IF",
    "Neftali Soto": "1B",
    "Nick Travieso": "P",
    "Javier Vaz": "IF",
}.items()}


def display_name(value: str) -> str:
    value = re.sub(r'\s+"[^"]+"', "", str(value or "").strip())
    if "," in value:
        last, first = [s.strip() for s in value.split(",", 1)]
        value = f"{first} {last}"
    return " ".join(part.capitalize() if part.isupper() else part for part in value.split())


def stable_id(team: str, name: str) -> int:
    raw = int(hashlib.sha1(f"{team}:{norm(name)}".encode()).hexdigest()[:8], 16)
    return -(raw % 800_000_000 + 100_000_000)


def collect_identity_rows(old_rosters: dict) -> list[dict]:
    rows = []
    for team in old_rosters.values():
        for p in team.get("players") or []:
            rows.append({"id": p.get("id"), "name": p.get("name"), "position": p.get("position"), "jersey": p.get("jersey")})
    for filename in (
        "lbprc_2025_hitting.json",
        "lbprc_2025_pitching.json",
        "mlb_2026_hitting.json",
        "mlb_2026_pitching.json",
        "milb_2026_hitting.json",
        "milb_2026_pitching.json",
        "mexico_2026_hitting.json",
        "mexico_2026_pitching.json",
    ):
        for p in load(filename):
            rows.append(
                {
                    "id": p.get("playerId") or p.get("mlbam_id"),
                    "name": p.get("playerFullName") or p.get("playerName") or p.get("name"),
                    "position": p.get("position") or ("P" if "pitching" in filename else None),
                    "jersey": "",
                }
            )
    for filename, key, id_key in (
        ("advance_opposing_hitters.json", "hitters", "mlb_id"),
        ("advance_opposing_pitchers.json", "pitchers", "mlb_id"),
        ("pregame_tth_pop.json", "players", "player_id"),
        ("pregame_pitcher_arsenals.json", "pitchers", "player_id"),
    ):
        for p in load(filename).get(key) or []:
            rows.append({"id": p.get(id_key), "name": p.get("name"), "position": p.get("position"), "jersey": p.get("jersey")})
    return [r for r in rows if r.get("id") and r.get("name")]


def identity_index(rows: list[dict]):
    by_name = defaultdict(list)
    for row in rows:
        by_name[norm(row["name"])].append(row)
    return by_name


ALIASES = {
    norm("Yamvier Carrier"): norm("Yamvier Carrero"),
    norm("Brian Navaretto"): norm("Brian Navarreto"),
    norm("Johneswy Fargas"): norm("Johneshwy Fargas"),
    norm("Carlos Espinosa"): norm("Carlos Espinosa"),
    norm("Jose Joe Joe Rodriguez"): norm("Jose Rodriguez"),
    norm("Luis Quinonez"): norm("Luis Quinones"),
    norm("Roberto Merced"): norm("Robbie Merced"),
    norm("Daniel Ortiz"): norm("Danny Ortiz"),
    norm("John Montes"): norm("John Montes"),
    norm("Chris Arroyo"): norm("Chris Arroyo"),
    norm("Roberto Enriquez"): norm("Roby Enriquez"),
}


def resolve_identity(name: str, position: str, by_name: dict, used_ids: set[int], team: str):
    key = norm(name)
    candidates = by_name.get(key) or by_name.get(ALIASES.get(key, "")) or []
    if not candidates:
        scored = []
        for other, values in by_name.items():
            score = SequenceMatcher(None, key, other).ratio()
            if score >= 0.94:
                scored.extend((score, row) for row in values)
        if scored:
            best = max(score for score, _ in scored)
            candidates = [row for score, row in scored if score == best]
    candidates = sorted(candidates, key=lambda r: (r.get("position") == position, bool(r.get("id"))), reverse=True)
    for row in candidates:
        pid = int(row["id"])
        if pid not in used_ids:
            return pid, row
    return stable_id(team, name), {}


def parse_base_rosters() -> dict[str, list[dict]]:
    rosters: dict[str, list[dict]] = {team: [] for team in TEAM_META}
    ws = load_workbook(SANTURCE_BOOK, data_only=True, read_only=True).active
    section = None
    seen = set()
    for row_num in range(4, 85):
        number, raw_name, bt = ws.cell(row_num, 1).value, ws.cell(row_num, 2).value, ws.cell(row_num, 3).value
        label = str(raw_name or "").upper()
        if "PITCHERS" in label:
            section = "P"
            continue
        if "CATCHERS" in label:
            section = "C"
            continue
        if "INFIELDERS" in label:
            section = "IF"
            continue
        if "OUTFIELDERS" in label:
            section = "OF"
            continue
        if label in {"RESERVE LIST", "IMPORT PLAYERS"}:
            section = "RES"
            continue
        if not raw_name or str(raw_name).upper() in {"STAFF", "2026 SELECCIONES DE SORTEO"}:
            continue
        name = display_name(str(raw_name))
        key = norm(name)
        if key in seen:
            continue
        seen.add(key)
        rosters["SAN"].append({"name": name, "position": section, "jersey": str(number or ""), "bt": bt})

    wb = load_workbook(LBPRC_BOOK, data_only=True, read_only=True)
    for sheet_name, team in SHEET_TEAM.items():
        ws = wb[sheet_name]
        for raw_name, raw_pos, _org in ws.iter_rows(min_row=2, values_only=True):
            if not raw_name:
                continue
            pos = str(raw_pos or "").upper()
            rosters[team].append(
                {
                    "name": display_name(str(raw_name)),
                    "position": "P" if pos in PITCHER_POS else (pos or "IF"),
                    "jersey": "",
                    "bt": None,
                }
            )

    additions = {
        "SJU": [("Yuta Nakami", "OF"), ("Keishin Saito", "P"), ("Haruse Narita", "P"), ("Freddy Cabrera", "P")],
        "PON": [("Mario Zabala", "P")],
    }
    moved = {norm(name) for values in additions.values() for name, _ in values}
    for team in rosters:
        rosters[team] = [p for p in rosters[team] if norm(p["name"]) not in moved]
    for team, values in additions.items():
        rosters[team].extend({"name": name, "position": pos, "jersey": "", "bt": None} for name, pos in values)
    return rosters


def materialize_rosters(base: dict[str, list[dict]], identities: list[dict], old_rosters: dict):
    by_name = identity_index(identities)
    old_by_id = {int(p["id"]): p for t in old_rosters.values() for p in t.get("players") or [] if p.get("id")}
    used_ids: set[int] = set()
    output = {}
    unresolved = []
    for team, entries in base.items():
        team_id, team_name = TEAM_META[team]
        players = []
        for item in entries:
            pos = item["position"]
            pid, prior = resolve_identity(item["name"], pos, by_name, used_ids, team)
            used_ids.add(pid)
            old = old_by_id.get(pid, {})
            if pos == "RES":
                pos = (
                    SANTURCE_RESERVE_POSITIONS.get(norm(item["name"]))
                    or old.get("position")
                    or prior.get("position")
                    or "IF"
                )
            pos = POSITION_NAMES.get(str(pos).upper(), str(pos).upper())
            if pos in PITCHER_POS:
                pos = "P"
            player = {
                "id": pid,
                "name": item["name"],
                "position": pos,
                "jersey": item["jersey"] or old.get("jersey") or prior.get("jersey") or "",
                "status": "Active",
            }
            players.append(player)
            if pid < 0:
                unresolved.append({"team": team, "name": item["name"], "position": pos, "synthetic_id": pid})
        output[str(team_id)] = {"teamId": team_id, "teamName": team_name, "teamAbbrev": team, "players": players}
    return output, unresolved


def roster_maps(rosters: dict):
    all_players, opposing, hitters, pitchers = [], [], [], []
    for team in rosters.values():
        for p in team["players"]:
            row = {**p, "team": team["teamAbbrev"], "team_name": team["teamName"]}
            all_players.append(row)
            if row["team"] != "SAN":
                opposing.append(row)
                (pitchers if row["position"] == "P" else hitters).append(row)
    return all_players, opposing, hitters, pitchers


def match_existing(player: dict, rows: list[dict], id_keys: tuple[str, ...]):
    for row in rows:
        for key in id_keys:
            if player["id"] > 0 and row.get(key) is not None and int(row[key]) == player["id"]:
                return copy.deepcopy(row)
    key = norm(player["name"])
    exact = [row for row in rows if norm(row.get("name") or row.get("playerFullName") or "") == key]
    return copy.deepcopy(exact[0]) if exact else None


def rebuild_advance(rosters: dict, hitters: list[dict], pitchers: list[dict]):
    hitter_doc = load("advance_opposing_hitters.json")
    old_hitters = hitter_doc.get("hitters") or []
    new_hitters = []
    for p in hitters:
        row = match_existing(p, old_hitters, ("mlb_id",))
        if row is None:
            row = {
                "synergy_id": None,
                "mlb_id": p["id"] if p["id"] > 0 else None,
                "name": p["name"],
                "bats": None,
                "overall": {},
                "by_pitch_type": [],
                "vs_l_by_pitch_type": [],
                "vs_r_by_pitch_type": [],
                "risp_by_pitch_type": [],
                "spray_points": [],
                "spray_splits": {},
                "summer_line": None,
            }
        row.update({"name": p["name"], "team": p["team"]})
        new_hitters.append(row)
    hitter_doc["updated"] = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    hitter_doc["source"] = "2026-27 LBPRC spreadsheet roster gate · retained historical scouting samples"
    hitter_doc.setdefault("season_meta", {})["roster_source"] = "2026-27 authoritative workbooks plus Oct. 2026 transactions"
    hitter_doc["hitters"] = sorted(new_hitters, key=lambda r: (r["team"], r["name"]))
    hitter_doc["select_all"] = None
    hitter_doc.setdefault("sample", {})["hitters"] = len(new_hitters)
    hitter_doc["notes"] = [
        "Roster gate replaced with the 2026-27 LBPRC workbooks plus the three supplied posts.",
        "New roster players without an existing scouting sample are included as empty records.",
    ]
    write("advance_opposing_hitters.json", hitter_doc)

    pitcher_doc = load("advance_opposing_pitchers.json")
    old_pitchers = pitcher_doc.get("pitchers") or []
    new_pitchers = []
    for p in pitchers:
        row = match_existing(p, old_pitchers, ("mlb_id",))
        if row is None:
            row = {
                "synergy_id": None,
                "mlb_id": p["id"] if p["id"] > 0 else None,
                "name": p["name"],
                "pitches": 0,
                "count_mix": {},
                "platoon": {},
                "leagues": {},
                "years": [],
                "source": "2026-27 roster stub; no scouting sample",
                "summer_line": None,
            }
        row.update({"name": p["name"], "team": p["team"]})
        new_pitchers.append(row)
    pitcher_doc["updated"] = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    pitcher_doc.setdefault("season_meta", {})["roster_source"] = "2026-27 authoritative workbooks plus Oct. 2026 transactions"
    pitcher_doc["pitchers"] = sorted(new_pitchers, key=lambda r: (r["team"], r["name"]))
    pitcher_doc["empty"] = [p["name"] for p in new_pitchers if not p.get("pitches")]
    write("advance_opposing_pitchers.json", pitcher_doc)


def rebuild_simple_caches(rosters: dict, all_players: list[dict], hitters: list[dict], pitchers: list[dict]):
    valid_all = {p["id"]: p for p in all_players}
    valid_hitters = {p["id"]: p for p in hitters}
    valid_pitchers = {p["id"]: p for p in pitchers}
    san_hitters = {p["id"]: p for p in all_players if p["team"] == "SAN" and p["position"] != "P"}
    san_pitchers = {p["id"]: p for p in all_players if p["team"] == "SAN" and p["position"] == "P"}

    lineup = [r for r in load("pregame_lineup_spots.json") if int(r.get("player_id") or 0) in san_hitters]
    write("pregame_lineup_spots.json", lineup)

    pvb = []
    for r in load("pregame_pitcher_vs_batter.json"):
        pit, hit = int(r.get("pitcher_id") or 0), int(r.get("hitter_id") or 0)
        if pit in san_pitchers and hit in valid_hitters:
            r["hitter_team"] = valid_hitters[hit]["team"]
            r["hitter_team_name"] = valid_hitters[hit]["team_name"]
            pvb.append(r)
    write("pregame_pitcher_vs_batter.json", pvb)

    hvp = []
    for r in load("pregame_hitter_vs_pitcher.json"):
        hit, pit = int(r.get("hitter_id") or 0), int(r.get("pitcher_id") or 0)
        if hit in san_hitters and pit in valid_pitchers:
            r["pitcher_team"] = valid_pitchers[pit]["team"]
            r["pitcher_team_name"] = valid_pitchers[pit]["team_name"]
            hvp.append(r)
    write("pregame_hitter_vs_pitcher.json", hvp)

    br = []
    for r in load("pregame_baserunning_bunting.json"):
        pid = int(r.get("player_id") or 0)
        if pid in valid_hitters:
            r["team"] = valid_hitters[pid]["team"]
            br.append(r)
    write("pregame_baserunning_bunting.json", br)

    spray = load("pregame_spray_charts.json")
    for key in ("mlb", "aaa_proxy", "synergy"):
        rows = []
        for r in spray.get(key) or []:
            pid = int(r.get("player_id") or r.get("mlb_id") or 0)
            if pid in valid_hitters:
                r["team"] = valid_hitters[pid]["team"]
                rows.append(r)
        if key in spray:
            spray[key] = rows
    write("pregame_spray_charts.json", spray)

    tth = load("pregame_tth_pop.json")
    old_tth = tth.get("players") or []
    tth_rows = []
    for p in all_players:
        row = match_existing(p, old_tth, ("player_id",)) or {}
        row.update(
            {
                "player_id": p["id"],
                "name": p["name"],
                "team": p["team"],
                "position": p["position"],
                "jersey": p["jersey"],
                "status": "Active",
                "time_to_home": row.get("time_to_home"),
                "catcher_pop_time": row.get("catcher_pop_time"),
            }
        )
        tth_rows.append(row)
    tth["meta"]["season"] = "2026-27 LBPRC rosters"
    tth["players"] = tth_rows
    write("pregame_tth_pop.json", tth)

    stuff = load("pregame_opposing_stuff.json")
    stuff_rows = []
    for p in pitchers:
        row = match_existing(p, stuff, ("player_id",)) or {
            "mlb_debut": None,
            "birth_country": None,
            "stuff_plus": None,
            "pitches": None,
            "arsenal": [],
            "source": "unavailable",
            "level_hint": None,
        }
        row.update({"name": p["name"], "player_id": p["id"], "team": p["team"]})
        stuff_rows.append(row)
    write("pregame_opposing_stuff.json", stuff_rows)

    arsenal_doc = load("pregame_pitcher_arsenals.json")
    old_arsenals = arsenal_doc.get("pitchers") or []
    arsenal_rows = []
    for p in pitchers:
        row = match_existing(p, old_arsenals, ("player_id",)) or {
            "throws": None,
            "mlb_debut": None,
            "level_hint": None,
            "stuff_plus": None,
            "pitches": 0,
            "arsenal": [],
            "usage_by_count": {},
            "count_mix": {},
            "platoon": {},
            "ps_season": None,
            "ps_source": None,
            "statcast_source": None,
            "statcast_seasons": [],
            "statcast_pitches": None,
            "view": "hitter",
            "note": "2026-27 roster stub; no arsenal sample",
        }
        row.update({"name": p["name"], "player_id": p["id"], "team": p["team"]})
        arsenal_rows.append(row)
    arsenal_doc["pitchers"] = arsenal_rows
    arsenal_doc["updated"] = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    write("pregame_pitcher_arsenals.json", arsenal_doc)


def main():
    old_rosters = load("lbprc_2025_rosters.json")
    base = parse_base_rosters()
    rosters, unresolved = materialize_rosters(base, collect_identity_rows(old_rosters), old_rosters)
    write("lbprc_2025_rosters.json", rosters)
    all_players, _opposing, hitters, pitchers = roster_maps(rosters)
    rebuild_advance(rosters, hitters, pitchers)
    rebuild_simple_caches(rosters, all_players, hitters, pitchers)
    audit = {
        "updated": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "season": "2026-27",
        "base_sources": [str(SANTURCE_BOOK), str(LBPRC_BOOK)],
        "transactions": TRANSACTIONS,
        "team_counts": {team["teamAbbrev"]: len(team["players"]) for team in rosters.values()},
        "unmatched_players": unresolved,
        "note": "Synthetic negative IDs keep unmatched spreadsheet/post additions visible without inventing MLBAM IDs.",
    }
    write("roster_2026_27_audit.json", audit)
    print("team counts", audit["team_counts"])
    print("synthetic ids", len(unresolved))
    for row in unresolved:
        print(" ", row)


if __name__ == "__main__":
    main()
